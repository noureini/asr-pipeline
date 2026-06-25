"""Extract structured wiki from the MMI survey instrument docx.

Produces:
  data/survey_wiki/
    instrument_full.txt       — entire instrument plaintext
    consent_script.txt        — verbatim consent intro (Section 1)
    consent_keywords.txt      — characteristic English keywords/phrases
                                that signal consent reading (for
                                ENUMERATOR verification against
                                Bengali ASR output)
    domain_vocabulary.md      — survey domain terms grouped by category
    sections.json             — section markers + question text per
                                section (for chunk-level section
                                tagging downstream)

These artifacts feed:
  - cleanup_titu.py (vocabulary as a domain anchor in the prompt)
  - verify_speaker_attribution.py (consent keywords to validate
    the heuristic ENUMERATOR detection in assemble_transcript.py)
  - The v2 architecture (instrument-aware diarization, section
    identification, accumulating-knowledge survey wiki).

Usage:
  uv run python scripts/build_survey_wiki.py \\
      --instrument "for testing (Noureini)/survey_instrument.docx"
"""
from __future__ import annotations
import argparse, json, re, sys
from pathlib import Path
from collections import defaultdict

REPO = Path(__file__).resolve().parents[1]


def extract_all(p: Path) -> tuple[list[str], list[list[list[str]]]]:
    """Return (paragraphs, tables_as_lists_of_rows_of_cell_texts)."""
    import docx
    d = docx.Document(str(p))
    paras = []
    for para in d.paragraphs:
        t = (para.text or "").strip()
        if t:
            paras.append(t)
    tables = []
    for t in d.tables:
        rows = []
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            rows.append(cells)
        tables.append(rows)
    return paras, tables


def extract_consent_script(paras: list[str]) -> str:
    """Pull the verbatim consent text — paragraphs between
    'SECTION 1' and 'SECTION 2' that are quoted speech.
    """
    in_section1 = False
    out = []
    for p in paras:
        if "SECTION 1" in p.upper():
            in_section1 = True
            continue
        if in_section1 and (
                "SECTION 2" in p.upper()
                or p.upper().startswith("S1.")):
            # S1.x are interviewer-private instructions, not script
            if "SECTION 2" in p.upper():
                break
            continue
        if in_section1:
            # Keep quoted blocks and the verbatim instruction lead-ins
            if p.startswith("“") or p.startswith('"') \
                    or "READ OUT" in p.upper():
                out.append(p)
            elif out and not p.startswith("IF") \
                    and not p.startswith("S1."):
                # continuation of a script block
                out.append(p)
    return "\n\n".join(out)


def extract_consent_keywords(consent_text: str) -> list[str]:
    """Pick characteristic English keywords/phrases that strongly
    signal consent reading — words a respondent would NOT casually
    use, but appear when the interviewer reads the script. These
    survive the interviewer's on-the-fly Bengali translation by
    appearing as code-switched English in the Bengali transcript.
    """
    candidates = [
        # Org / formal names (always English in script)
        "Data Analysis", "Technical Assistance",
        # Formal nouns that get spoken in English
        "audio-record", "audio record", "audio recording",
        "interview", "research", "confidential",
        "voluntary", "consent",
        # Numbers & duration phrases
        "1 hour", "one hour",
        # Section-1-only verbs
        "participate", "questions",
        # Roles
        "researcher", "interviewer", "respondent",
    ]
    # Keep only ones that actually appear in our consent_text
    found = [k for k in candidates if k.lower()
             in consent_text.lower()]
    return found


# Generic English stopwords + survey-machinery words to exclude from the
# extracted domain terms (these recur in any instrument and aren't domain
# vocabulary).
_STOP = set("""
a an the and or but if then else of to in for on at by with from into onto
as is are was were be been being do does did have has had will would shall
should can could may might must not no nor this that these those it its it's
you your yours we our ours they them their he she his her him i me my mine
which who whom whose what when where why how all any both each few more most
other some such only own same so than too very s t just don now also about
above below up down out off over under again further once here there
section part question questions answer answers respondent interviewer
enumerator read out note instruction instructions code codes option options
please ask record write circle select choose following above below page
yes maybe other specify text date time name number
""".split())

# Survey/section header tokens to drop from acronym/entity extraction.
_HEADER_ACRO = {"SECTION", "PART", "READ", "OUT", "NOTE", "IF", "AND", "OR",
                "THE", "FOR", "YES", "NO", "ID", "Q", "A", "S"}


def extract_vocabulary(paras: list[str],
                        tables: list[list[list[str]]]
                        ) -> dict[str, list[str]]:
    """Extract domain vocabulary GENERICALLY from the instrument text —
    no hardcoded, survey-specific terms. Works for any instrument docx.

    Three faithful sources, all drawn from the actual document:
      named_entities  — acronyms (NGO, BRAC, VGD...) + Title-Case
                        multiword names (Krishi Bank, World Vision...)
      answer_options  — short enumerated choices from question tables
                        (the controlled vocabulary the survey uses)
      domain_terms    — frequent non-stopword content words (the terms
                        that characterise THIS survey's subject matter)
    """
    import re as _re
    from collections import Counter

    full = list(paras)
    for tbl in tables:
        for row in tbl:
            full.append(" ".join(c for c in row if c.strip()))
    text = "\n".join(full)

    ne: set[str] = set()
    # Acronyms: 2-6 uppercase letters/digits (NGO, BRAC, VGD, SSC, VGF)
    for m in _re.findall(r"\b[A-Z][A-Z0-9]{1,5}\b", text):
        if m not in _HEADER_ACRO:
            ne.add(m)
    # Title-Case multiword entities: "Krishi Bank", "World Vision"
    for m in _re.findall(
            r"\b([A-Z][a-z]+(?:\s+[A-Z][a-z]+){1,3})\b", text):
        head = m.split()[0].upper()
        if head not in _HEADER_ACRO:
            ne.add(m)

    # Answer options: cells/lines that are short enumerated choices.
    opts: set[str] = set()
    for tbl in tables:
        for row in tbl:
            for cell in row:
                # split a cell that packs options like "1. Islam 2. Hindu"
                for opt in _re.split(
                        r"\s*(?:\d+[\.\)]|[a-z][\.\)]|/|;|,)\s*", cell):
                    opt = opt.strip(" .:–-")
                    nw = len(opt.split())
                    if (opt and 1 <= nw <= 3 and opt[0].isalpha()
                            and opt.lower() not in _STOP
                            and not opt.isupper()):
                        opts.add(opt)

    # Domain terms: frequent lowercase content words actually in the doc.
    words = _re.findall(r"\b[a-z][a-z\-]{3,}\b", text.lower())
    freq = Counter(w for w in words if w not in _STOP)
    domain = {w for w, c in freq.most_common(80) if c >= 3}

    return {
        "named_entities": sorted(ne),
        "answer_options": sorted(opts),
        "domain_terms": sorted(domain),
    }


def extract_sections(paras: list[str],
                      tables: list[list[list[str]]]) -> list[dict]:
    """Identify section markers and any verbatim question text we
    can attach to each section. Output is a list of
    {section_id, title, marker_text, questions[]}.
    """
    sections: list[dict] = []
    section_markers = [
        ("SECTION 1", "Introduction & Consent"),
        ("SECTION 2", "Respondent Basic Information"),
        ("SECTION 3", "Interview Protocol"),
        ("PART I", "Part I — Initial mental model"),
        ("PART II-A", "Part II-A — Floods specifically"),
        ("PART II-B", "Part II-B — Floods continued"),
        ("TRANSITION FROM PART I TO PART II",
         "Bridge: Part I → Part II"),
        ("TRANSITION FROM PART I TO PART III",
         "Bridge: Part I → Part III (no floods mentioned)"),
        ("PART III", "Part III — Adaptation & coping"),
        ("PART IV", "Part IV — Future outlook"),
        ("CONCLUDING REMARKS", "Closing"),
    ]
    for marker, title in section_markers:
        if any(marker.upper() in p.upper() for p in paras):
            sections.append({
                "section_id": marker.replace(" ", "_"),
                "title": title,
                "marker_text": marker,
                "questions": [],
            })
    # Attach extracted question text from tables where possible
    for tbl in tables:
        for row in tbl:
            for cell in row:
                if re.match(r"^S\d", cell.strip()):
                    qid = cell.split(".")[0].strip()
                    sec = qid[:2]  # S0, S1, S2, S3 ...
                    target = None
                    for s in sections:
                        if s["marker_text"].startswith(
                                f"SECTION {sec[1]}"):
                            target = s
                            break
                    if target is not None:
                        text = cell.strip()
                        if text not in target["questions"]:
                            target["questions"].append(text)
    return sections


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument(
        "--instrument", type=Path,
        default=Path("for testing (Noureini)")
                / "survey_instrument.docx")
    ap.add_argument("--out-dir", type=Path,
                    default=REPO / "data" / "survey_wiki")
    args = ap.parse_args()

    if not args.instrument.exists():
        sys.exit(f"missing instrument: {args.instrument}")
    args.out_dir.mkdir(parents=True, exist_ok=True)

    paras, tables = extract_all(args.instrument)
    print(f"loaded instrument: {len(paras)} paragraphs, "
          f"{len(tables)} tables")

    # 1. instrument_full.txt — paragraphs + flattened tables
    full = list(paras)
    for ti, tbl in enumerate(tables):
        full.append(f"\n=== Table {ti} ===")
        for row in tbl:
            non_empty = [c for c in row if c.strip()]
            if non_empty:
                full.append(" | ".join(non_empty))
    (args.out_dir / "instrument_full.txt").write_text(
        "\n".join(full), encoding="utf-8")
    print(f"  -> instrument_full.txt "
          f"({len('\\n'.join(full))} chars)")

    # 2. consent_script.txt
    consent = extract_consent_script(paras)
    (args.out_dir / "consent_script.txt").write_text(
        consent, encoding="utf-8")
    print(f"  -> consent_script.txt ({len(consent)} chars)")

    # 3. consent_keywords.txt
    keywords = extract_consent_keywords(consent)
    (args.out_dir / "consent_keywords.txt").write_text(
        "\n".join(keywords) + "\n", encoding="utf-8")
    print(f"  -> consent_keywords.txt ({len(keywords)} keywords)")

    # 4. domain_vocabulary.md
    vocab = extract_vocabulary(paras, tables)
    lines = ["# Survey domain vocabulary",
             "",
             f"GENERICALLY EXTRACTED from this instrument "
             f"({args.instrument.name}) — named entities, enumerated "
             "answer options, and frequent domain terms found in the "
             "document text. No survey-specific terms are hardcoded, so "
             "this reflects ONLY what is actually in the instrument. "
             "Used as a domain anchor in the ASR fusion prompt.",
             ""]
    for cat, terms in vocab.items():
        lines.append(f"## {cat}  ({len(terms)})")
        lines.append("")
        lines.append(", ".join(terms) if terms else "(none found)")
        lines.append("")
    (args.out_dir / "domain_vocabulary.md").write_text(
        "\n".join(lines), encoding="utf-8")
    n_terms = sum(len(v) for v in vocab.values())
    print(f"  -> domain_vocabulary.md ({n_terms} terms: "
          f"{len(vocab['named_entities'])} entities, "
          f"{len(vocab['answer_options'])} options, "
          f"{len(vocab['domain_terms'])} domain terms)")

    # 5. sections.json
    sections = extract_sections(paras, tables)
    (args.out_dir / "sections.json").write_text(
        json.dumps(sections, ensure_ascii=False, indent=2),
        encoding="utf-8")
    print(f"  -> sections.json ({len(sections)} sections, "
          f"{sum(len(s['questions']) for s in sections)} "
          f"questions tagged)")

    # 6. Placeholder for accumulating knowledge (per v2 architecture)
    for ph in ("proper_nouns.txt", "corrections.md",
                "common_phrases.md"):
        f = args.out_dir / ph
        if not f.exists():
            f.write_text(
                f"# {ph}\n\n# (Grows as interviews are processed)\n",
                encoding="utf-8")
    print(f"  -> proper_nouns.txt, corrections.md, "
          f"common_phrases.md (placeholders for v2)")

    print(f"\nsurvey wiki ready at: {args.out_dir}")


if __name__ == "__main__":
    main()
